"""Scrap statistics over a date range (the Scrap statistics page)."""
import datetime as dt
import io

from openpyxl import load_workbook

from app.database import SessionLocal
from app.models import Reading, Station

from test_api import login

UTC = dt.timezone.utc
RANGE = {"from": "2026-09-29", "to": "2026-09-30", "tz": "UTC"}


def _seed(in_production=True):
    """Readings of one station; returns their ids in order."""
    db = SessionLocal()
    try:
        cam = Station(name="Cam1", sources={})
        db.add(cam)
        db.flush()
        rows = []

        def r(day, hour, job, p, f, month=9):
            row = Reading(station_id=cam.id, job_name=job, total_pass=p, total_fail=f, in_production=in_production,
                          created_at=dt.datetime(2026, month, day, hour, tzinfo=UTC))
            db.add(row)
            rows.append(row)

        r(28, 22, "A", 100, 10)  # before the range: baseline only
        r(29, 8, "A", 190, 20)   # +90/+10 on the 29th
        r(29, 9, "B", 5, 5)      # job change: new baseline, not counted
        r(30, 8, "B", 45, 5)     # +40/+0 on the 30th
        r(30, 9, "B", 45, 8)     # +0/+3
        r(2, 8, "B", 999, 999, month=10)  # after the range
        db.commit()
        return [x.id for x in rows]
    finally:
        db.close()


def _set(ids, **values):
    db = SessionLocal()
    try:
        for i in ids:
            r = db.get(Reading, i)
            for k, v in values.items():
                setattr(r, k, v)
        db.commit()
    finally:
        db.close()


def test_scrap_range_per_camera_job_day(client):
    login(client)
    _seed()
    s = client.get("/api/data/scrap", params=RANGE).json()
    assert s["overall"] == {"pass": 130, "fail": 13, "total": 143, "scrap_rate": round(13 / 143, 4)}
    assert [(c["station"], c["total"]) for c in s["per_station"]] == [("Cam1", 143)]
    assert [(j["job"], j["pass"], j["fail"]) for j in s["per_job"]] == [("A", 90, 10), ("B", 40, 3)]
    assert [(d["day"], d["pass"], d["fail"]) for d in s["per_day"]] == [
        ("2026-09-29", 90, 10), ("2026-09-30", 40, 3)]


def test_scrap_days_follow_time_zone(client):
    login(client)
    _seed()
    # The 30th in Tokyo (UTC+9) runs from 15:00 UTC on the 29th, so the job B
    # reading at 09:00 UTC on the 29th is the baseline and both readings on
    # the 30th (08:00 and 09:00 UTC) fall inside it.
    s = client.get("/api/data/scrap", params={"from": "2026-09-30", "to": "2026-09-30", "tz": "Asia/Tokyo"}).json()
    assert [(d["day"], d["total"]) for d in s["per_day"]] == [("2026-09-30", 43)]


def test_excluded_readings_are_left_out(client):
    login(client)
    ids = _seed()
    # one reading by itself...
    assert client.patch(f"/api/data/readings/{ids[1]}", json={"excluded": True}).json()["excluded"] is True
    s = client.get("/api/data/scrap", params=RANGE).json()
    assert (s["overall"]["pass"], s["overall"]["fail"]) == (40, 3)
    assert s["left_out"]["excluded"] == {"pass": 90, "fail": 10, "total": 100, "scrap_rate": 0.1, "readings": 1}
    # the next reading still only counts its own growth
    assert [(j["job"], j["total"]) for j in s["per_job"]] == [("B", 43)]
    # ...and a whole period, then taken back
    period = {"start": "2026-09-30T00:00:00Z", "end": "2026-10-01T00:00:00Z"}
    assert client.post("/api/data/readings/exclude", json={**period, "excluded": True}).json() == {"changed": 2}
    assert client.get("/api/data/scrap", params=RANGE).json()["overall"]["total"] == 0
    assert len(client.get("/api/data/readings", params={"excluded": "true"}).json()) == 3
    client.post("/api/data/readings/exclude", json={**period, "excluded": False, "station_id": 1})
    client.patch(f"/api/data/readings/{ids[1]}", json={"excluded": False})
    assert client.get("/api/data/scrap", params=RANGE).json()["overall"]["total"] == 143


def test_not_in_production_is_left_out(client):
    login(client)
    ids = _seed()
    _set([ids[4]], in_production=False)  # the +3 fail while idle or stopped
    s = client.get("/api/data/scrap", params=RANGE).json()
    assert (s["overall"]["pass"], s["overall"]["fail"]) == (130, 10)
    assert (s["left_out"]["not_in_production"]["pass"], s["left_out"]["not_in_production"]["fail"]) == (0, 3)


def test_older_readings_use_the_idle_rule(client):
    login(client)
    _seed(in_production=None)
    # default idle timeout is 30 min: the +3 fail at 09:00 on the 30th came an
    # hour after the last pass increase (08:00), so the camera was idle
    s = client.get("/api/data/scrap", params=RANGE).json()
    assert (s["overall"]["pass"], s["overall"]["fail"]) == (130, 10)
    assert s["left_out"]["not_in_production"]["fail"] == 3


def test_scrap_xlsx_matches_screen(client):
    login(client)
    _seed()
    resp = client.get("/api/data/scrap.xlsx", params=RANGE)
    assert resp.status_code == 200
    assert "scrap_2026-09-29_2026-09-30.xlsx" in resp.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Overall", "Per station", "Per job", "Per day"]
    assert [c.value for c in wb["Overall"][2]] == ["Counted, 2026-09-29 to 2026-09-30", 130, 13, 143, round(13 / 143, 4)]
    assert wb["Overall"]["A3"].value == "Left out: excluded readings"
    assert [c.value for c in wb["Per job"][3]][:5] == ["Cam1", "B", 40, 3, 43]
    assert wb["Per day"]["A2"].value.date() == dt.date(2026, 9, 29)


def test_scrap_range_validation_and_permissions(client):
    login(client)
    ids = _seed()
    assert client.get("/api/data/scrap", params={"from": "2026-09-30", "to": "2026-09-29"}).status_code == 400
    client.post("/api/users", json={"username": "viewer", "password": "pw123456", "permissions": ["view_dashboard"]})
    client.post("/api/users", json={"username": "reader", "password": "pw123456", "permissions": ["view_data"]})
    client.get("/logout")
    login(client, "viewer", "pw123456")
    assert client.get("/api/data/scrap", params=RANGE).status_code == 403
    assert client.get("/scrap", follow_redirects=False).status_code != 200
    client.get("/logout")
    login(client, "reader", "pw123456")
    assert client.get("/scrap").status_code == 200
    # excluding needs its own permission
    assert client.patch(f"/api/data/readings/{ids[1]}", json={"excluded": True}).status_code == 403


def test_poller_records_production_state(client):
    login(client)
    did = client.post("/api/devices", json={
        "name": "Sim", "host": "sim", "port": 0, "protocol": "simulator",
        "protocol_config": {"jobs": ["J"], "parts_per_poll": 10, "reset_every": 0, "job_change_every": 0},
        "create_station": True,
    }).json()["id"]
    client.post(f"/api/devices/{did}/poll")  # first sample: baseline, no pass increase yet
    client.post(f"/api/devices/{did}/poll")  # pass went up: in production
    states = [r["in_production"] for r in client.get("/api/data/readings").json()]
    assert states == [True, False]  # newest first
