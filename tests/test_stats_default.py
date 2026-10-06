"""Per-device statistics default: devices set to "exclude" stay out of the
overall figures but keep their own rows, and can be included again."""
import datetime as dt
import io
import re

from openpyxl import load_workbook

from app import commands
from app.database import SessionLocal
from app.models import ChatCommand, Device, Reading

from test_api import login

UTC = dt.timezone.utc
RANGE = {"from": "2026-09-29", "to": "2026-09-30", "tz": "UTC"}


def _seed():
    """Two devices with the same readings: Main (include) and Test rig (exclude).
    Each makes +90/+10 on the 29th and +40/+3 on the 30th."""
    db = SessionLocal()
    try:
        ids = {}
        for name, default in (("Main", "include"), ("Test rig", "exclude")):
            dev = Device(name=name, host="sim", port=0, protocol="simulator", stats_default=default)
            db.add(dev)
            db.flush()
            rows = []
            for day, hour, p, f in ((28, 22, 100, 10), (29, 8, 190, 20), (30, 8, 230, 20), (30, 9, 230, 23)):
                r = Reading(device_id=dev.id, job_name="A", total_pass=p, total_fail=f, in_production=True,
                            created_at=dt.datetime(2026, 9, day, hour, tzinfo=UTC))
                db.add(r)
                rows.append(r)
            db.flush()
            ids[name] = (dev.id, [r.id for r in rows])
        db.commit()
        return ids
    finally:
        db.close()


def test_new_and_existing_devices_default_to_include(client):
    login(client)
    r = client.post("/api/devices", json={"name": "D", "host": "sim", "port": 0, "protocol": "simulator"})
    assert r.json()["stats_default"] == "include"
    assert client.patch(f"/api/devices/{r.json()['id']}", json={"stats_default": "exclude"}).json()[
        "stats_default"] == "exclude"
    assert client.patch(f"/api/devices/{r.json()['id']}", json={"stats_default": "maybe"}).status_code == 422


def test_excluded_device_stays_out_of_overall(client):
    login(client)
    _seed()
    s = client.get("/api/data/scrap", params=RANGE).json()
    assert (s["overall"]["pass"], s["overall"]["fail"]) == (130, 13)
    assert [(d["day"], d["total"]) for d in s["per_day"]] == [("2026-09-29", 100), ("2026-09-30", 43)]
    rows = {r["camera"]: r for r in s["per_camera"]}
    # still shown on its own, marked
    assert (rows["Test rig"]["total"], rows["Test rig"]["excluded_by_default"]) == (143, True)
    assert (rows["Test rig"]["counted_pass"], rows["Test rig"]["counted_fail"]) == (0, 0)
    assert rows["Main"]["excluded_by_default"] is False
    assert [j["excluded_by_default"] for j in s["per_job"]] == [False, True]
    assert s["left_out"]["excluded_devices"]["total"] == 143


def test_readings_of_excluded_device_can_be_included(client):
    login(client)
    ids = _seed()
    rig_id, rig_rows = ids["Test rig"]
    # one reading: the +90/+10 on the 29th
    r = client.patch(f"/api/data/readings/{rig_rows[1]}", json={"excluded": False}).json()
    assert r["included"] is True
    s = client.get("/api/data/scrap", params=RANGE).json()
    assert s["overall"]["total"] == 143 + 100
    # a period, then excluded again
    period = {"start": "2026-09-30T00:00:00Z", "end": "2026-10-01T00:00:00Z", "device_id": rig_id}
    client.post("/api/data/readings/exclude", json={**period, "excluded": False})
    assert client.get("/api/data/scrap", params=RANGE).json()["overall"]["total"] == 286
    client.post("/api/data/readings/exclude", json={**period, "excluded": True})
    s = client.get("/api/data/scrap", params=RANGE).json()
    assert s["overall"]["total"] == 243
    assert s["left_out"]["excluded"]["total"] == 43
    # the data log shows what is left out by the device default
    out = client.get("/api/data/readings", params={"excluded": "true", "device_id": rig_id}).json()
    assert {x["id"] for x in out} == {rig_rows[0], rig_rows[2], rig_rows[3]}
    assert all(x["device_excluded"] for x in out)
    # including a whole period on an "include" device does not mark its readings
    client.post("/api/data/readings/exclude", json={**period, "device_id": ids["Main"][0], "excluded": False})
    db = SessionLocal()
    try:
        assert not any(db.get(Reading, i).included for i in ids["Main"][1])
    finally:
        db.close()


def test_xlsx_marks_excluded_devices(client):
    login(client)
    _seed()
    wb = load_workbook(io.BytesIO(client.get("/api/data/scrap.xlsx", params=RANGE).content))
    assert wb.sheetnames == ["Overall", "Per device", "Per job", "Per day"]
    assert [c.value for c in wb["Overall"][2]][1:4] == [130, 13, 143]
    assert wb["Overall"]["A5"].value == "Left out: devices excluded by default"
    assert wb["Overall"]["D5"].value == 143
    assert [c.value for c in wb["Per device"][3]] == ["Test rig", 130, 13, 143, round(13 / 143, 4),
                                                      "No (excluded by default)"]
    assert wb["Per device"]["F2"].value == "Yes"


def test_status_command_totals_skip_excluded_devices(client):
    login(client)
    for name in ("Line 1", "Rig"):
        did = client.post("/api/devices", json={
            "name": name, "host": "sim", "port": 0, "protocol": "simulator",
            "stats_default": "exclude" if name == "Rig" else "include",
            "protocol_config": {"jobs": ["J1"], "parts_per_poll": 10, "fail_ratio": 0.2,
                                "reset_every": 0, "job_change_every": 0}}).json()["id"]
        client.post(f"/api/devices/{did}/production/start")
        for _ in range(3):
            client.post(f"/api/devices/{did}/poll")
    db = SessionLocal()
    try:
        for period in ("dashboard", "hours", "today"):
            cmd = ChatCommand(**{**commands.DEFAULT_STATUS, "keyword": "s" + period, "period": period,
                                 "group_ids": []})
            db.add(cmd)
            db.commit()
            reply = commands.render(db, cmd)
            assert re.search(r"Rig: .*\n.*scrap [^\n]*\(not in totals\)", reply), reply
            counts = [int(a) + int(b) for a, b in re.findall(r"OK (\d+) · NOK (\d+) · scrap", reply)]
            line1 = counts[0]
            assert counts[-1] == line1, (period, reply)  # total = Line 1 only
    finally:
        db.close()


def test_export_import_carries_the_setting(client):
    login(client)
    client.post("/api/devices", json={"name": "Rig", "host": "sim", "port": 0, "protocol": "simulator",
                                      "stats_default": "exclude"})
    data = client.get("/api/devices/export").json()
    assert data["cameras"][0]["stats_default"] == "exclude"
    # files from older versions have no setting: imported as include
    old = {"version": 1, "cameras": [{"name": "Old", "host": "sim", "port": 0, "protocol": "simulator"}]}
    assert client.post("/api/devices/import", json=old).json()["created"] == ["Old"]
    data["cameras"][0]["name"] = "Rig 2"
    client.post("/api/devices/import", json=data)
    got = {d["name"]: d["stats_default"] for d in client.get("/api/devices").json()}
    assert got == {"Old": "include", "Rig": "exclude", "Rig 2": "exclude"}
