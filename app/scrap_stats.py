"""Scrap statistics over a date range (the Scrap statistics page).

Parts are counted per station the same way as the station view's OK/NOK
chart (``parts``): a reading since 1.8 says which pieces it added to the
station's counters (ok_added / nok_added); on older readings they are the
difference of the running totals of two consecutive readings of the same job,
and a job change starts a new baseline. Scrap is fail / (pass + fail).

Parts are left out (and reported separately) when their reading
* was excluded by a user, or
* was taken while the station was not in production (idle or stopped).
  Readings logged before that was recorded fall back to the idle rule: no
  pass increase within the station's idle timeout.

A station set to "exclude" from the statistics (Station.stats_default) still
gets its own per-station and per-job rows, marked ``excluded_by_default``, but
its parts stay out of the overall figures and the per-day rows unless a user
included the reading (Reading.included). Those parts are reported under
left_out.excluded_stations. ``counted_pass`` / ``counted_fail`` on a station
row are its parts that are in the overall figures.

``station_parts`` walks one station's readings; app.oee uses it too.

``report`` adds what the page shows on top: the previous period of the same
length (the change of each figure) and the scrap alert thresholds (the days
above them are marked).

Days are calendar days in the viewer's time zone; the range is inclusive.
"""
from __future__ import annotations

import datetime as dt
import io
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from sqlalchemy.orm import Session

from . import production
from .models import Reading, Station


def zone(name: str | None) -> dt.tzinfo:
    try:
        return ZoneInfo(name) if name else dt.timezone.utc
    except (ZoneInfoNotFoundError, ValueError):
        return dt.timezone.utc


def _aware(t: dt.datetime) -> dt.datetime:
    return t.replace(tzinfo=dt.timezone.utc) if t.tzinfo is None else t


def _row(**keys) -> dict:
    return {**keys, "pass": 0, "fail": 0}


def _finish(rows: list[dict]) -> list[dict]:
    for r in rows:
        r["total"] = r["pass"] + r["fail"]
        # a correction (negative manual entry) can take a total to zero or below
        r["scrap_rate"] = round(min(max(r["fail"] / r["total"], 0.0), 1.0), 4) if r["total"] > 0 else 0.0
    return rows


def day_bounds(first_day: dt.date, last_day: dt.date, tz: dt.tzinfo) -> tuple[dt.datetime, dt.datetime]:
    """UTC start and (exclusive) end of an inclusive range of local days."""
    start = dt.datetime.combine(first_day, dt.time(), tz).astimezone(dt.timezone.utc)
    end = dt.datetime.combine(last_day + dt.timedelta(days=1), dt.time(), tz).astimezone(dt.timezone.utc)
    return start, end


def parts(r, prev) -> tuple[int, int]:
    """(OK, NOK) a reading counted; ``prev`` is the station's previous device
    reading (only used for readings from before 1.8)."""
    if r.manual:  # entered by hand: its own parts (negative for a correction)
        return r.raw_pass, r.raw_fail
    if r.ok_added is not None:
        return r.ok_added, r.nok_added or 0
    if prev is not None and prev.job_name == r.job_name:
        return max(r.total_pass - prev.total_pass, 0), max(r.total_fail - prev.total_fail, 0)
    return 0, 0


# the Reading columns ``parts`` needs
PART_COLUMNS = (Reading.job_name, Reading.total_pass, Reading.total_fail, Reading.manual,
                Reading.raw_pass, Reading.raw_fail, Reading.ok_added, Reading.nok_added)


def station_parts(db: Session, station: Station, start: dt.datetime, end: dt.datetime):
    """The station's readings in [start, end) with the parts each one counted.

    Yields dicts: id, t, prev_t (the previous reading's time, or None), job,
    ok, nok (parts since the previous reading of the same job; a manual
    entry's own parts, negative for a correction), excluded, included (by a
    user), in_production (recorded, or by the idle rule).
    """
    cols = (*PART_COLUMNS, Reading.id, Reading.created_at, Reading.excluded, Reading.in_production, Reading.included)
    timeout = dt.timedelta(minutes=max(1, station.idle_timeout_min or production.DEFAULT_IDLE_TIMEOUT_MIN))
    # start one idle timeout early so the idle rule knows the last pass
    # increase before the range
    lead = start - timeout
    q = db.query(*cols).filter(Reading.station_id == station.id)
    baseline = (q.filter(Reading.created_at < lead, Reading.manual.isnot(True))
                .order_by(Reading.created_at.desc(), Reading.id.desc()).first())
    rows = (
        q.filter(Reading.created_at >= lead, Reading.created_at < end)
        .order_by(Reading.created_at.asc(), Reading.id.asc())
        .yield_per(5000)
    )
    prev = baseline
    last_increase = None
    for r in rows:
        t = _aware(r.created_at)
        if r.manual:
            # entered by hand: its own parts, outside the devices' totals
            if t >= start:
                yield {"id": r.id, "t": t, "prev_t": None, "job": r.job_name, "ok": r.raw_pass, "nok": r.raw_fail,
                       "excluded": bool(r.excluded), "included": bool(r.included), "in_production": True,
                       "manual": True}
            continue
        d_ok, d_nok = parts(r, prev)
        prev_t = _aware(prev.created_at) if prev is not None else None
        prev = r
        if d_ok:
            last_increase = t
        if t < start:
            continue
        in_prod = (r.in_production if r.in_production is not None
                   else last_increase is not None and t - last_increase < timeout)
        yield {"id": r.id, "t": t, "prev_t": prev_t, "job": r.job_name, "ok": d_ok, "nok": d_nok,
               "excluded": bool(r.excluded), "included": bool(r.included), "in_production": bool(in_prod)}


def compute(db: Session, first_day: dt.date, last_day: dt.date, tz: dt.tzinfo) -> dict:
    start, end = day_bounds(first_day, last_day, tz)

    days = {}
    d = first_day
    while d <= last_day:
        days[d] = _row(day=d.isoformat())
        d += dt.timedelta(days=1)
    overall = _row()
    left_out = {"excluded": _row(readings=0), "not_in_production": _row(), "excluded_stations": _row()}
    per_station: dict[int, dict] = {}
    per_job: dict[tuple[int, str], dict] = {}
    # each station's own parts per day (also for a station excluded by default)
    station_days: dict[int, dict[dt.date, list[int]]] = {}
    # a range of up to two days is also counted per hour, for the trend
    # (every hour, so hours without production show as gaps)
    hours = None
    if (last_day - first_day).days < 2:
        hours, t = {}, start
        while t < end:
            hour = t.astimezone(tz).replace(minute=0, second=0, microsecond=0).isoformat()
            hours.setdefault(hour, _row(hour=hour))
            t += dt.timedelta(hours=1)

    for st in db.query(Station).order_by(Station.sort_order, Station.name).all():
        by_default = st.excluded_by_default
        srow = per_station[st.id] = _row(station_id=st.id, station=st.name, excluded_by_default=by_default,
                                         counted_pass=0, counted_fail=0)
        for r in station_parts(db, st, start, end):
            d_ok, d_nok, t = r["ok"], r["nok"], r["t"]
            if not (d_ok or d_nok):
                continue
            if r["excluded"]:
                targets = (left_out["excluded"],)
                left_out["excluded"]["readings"] += 1
            elif not r["in_production"]:
                targets = (left_out["not_in_production"],)
            else:
                job = r["job"] or "—"
                jrow = per_job.setdefault(
                    (st.id, job), _row(station=st.name, job=job, excluded_by_default=by_default))
                local = t.astimezone(tz)
                cell = station_days.setdefault(st.id, {}).setdefault(local.date(), [0, 0])
                cell[0] += d_ok
                cell[1] += d_nok
                if by_default and not r["included"]:
                    targets = (srow, jrow, left_out["excluded_stations"])
                else:
                    targets = (overall, srow, jrow, days[local.date()])
                    if hours is not None:
                        hour = local.replace(minute=0, second=0, microsecond=0).isoformat()
                        targets += (hours.setdefault(hour, _row(hour=hour)),)
                    srow["counted_pass"] += d_ok
                    srow["counted_fail"] += d_nok
            for target in targets:
                target["pass"] += d_ok
                target["fail"] += d_nok

    excluded_readings = (
        db.query(Reading.id)
        .filter(Reading.excluded.is_(True), Reading.created_at >= start, Reading.created_at < end)
        .count()
    )
    return {
        "from": first_day.isoformat(),
        "to": last_day.isoformat(),
        "overall": _finish([overall])[0],
        "per_station": _finish(list(per_station.values())),
        "per_job": _finish(sorted(per_job.values(), key=lambda r: (r["station"], r["job"]))),
        "per_day": _finish(list(days.values())),
        # per station: [[pass, fail] for each day of per_day]
        "station_days": {str(sid): [cells.get(d, [0, 0]) for d in days] for sid, cells in station_days.items()},
        "per_hour": _finish(sorted(hours.values(), key=lambda r: r["hour"])) if hours is not None else None,
        "left_out": {
            "excluded": {**_finish([left_out["excluded"]])[0], "readings": excluded_readings},
            "not_in_production": _finish([left_out["not_in_production"]])[0],
            "excluded_stations": _finish([left_out["excluded_stations"]])[0],
        },
    }


# the scrap rate the page marks as high when no scrap alert rule is set (the
# red level of the scrap badges)
DEFAULT_ALERT = 0.05


def alert_thresholds(db: Session) -> dict:
    """The scrap rates the enabled "Scrap rate" alert rules fire at: the
    lowest one for all stations ("overall", for the per-day figures) and per
    station (its own rules and its overrides of the rules for all stations).
    Without such rules the page uses DEFAULT_ALERT."""
    from .models import NotificationRule
    from .notifications import threshold_for

    rules = (db.query(NotificationRule)
             .filter(NotificationRule.enabled.is_(True), NotificationRule.condition == "scrap_rate").all())
    if not rules:
        return {"source": "default", "overall": DEFAULT_ALERT, "per_station": {}}
    general = [r.threshold for r in rules if r.station_id is None]
    per_station = {}
    for st in db.query(Station).all():
        values = [threshold_for(r, st) for r in rules if r.station_id in (None, st.id)]
        if values:
            per_station[str(st.id)] = min(values)
    overall = min(general) if general else min(r.threshold for r in rules)
    return {"source": "rules", "overall": overall, "per_station": per_station}


def report(db: Session, first_day: dt.date, last_day: dt.date, tz: dt.tzinfo) -> dict:
    """``compute`` for the page: with the previous period of the same length
    (for the change of each figure) and the alert thresholds."""
    stats = compute(db, first_day, last_day, tz)
    n = (last_day - first_day).days + 1
    prev = compute(db, first_day - dt.timedelta(days=n), first_day - dt.timedelta(days=1), tz)
    stats["previous"] = {
        "from": prev["from"], "to": prev["to"], "overall": prev["overall"],
        "per_station": {str(r["station_id"]): {k: r[k] for k in ("pass", "fail", "total", "scrap_rate")}
                        for r in prev["per_station"]},
        "per_job": [{k: r[k] for k in ("station", "job", "pass", "fail", "total", "scrap_rate")}
                    for r in prev["per_job"]],
    }
    alert = alert_thresholds(db)
    alert["days_above"] = [r["day"] for r in stats["per_day"]
                           if r["total"] > 0 and r["scrap_rate"] >= alert["overall"]]
    stats["alert"] = alert
    return stats


def to_xlsx(stats: dict) -> bytes:
    """One sheet per view, with the same figures as the screen."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    period = f"{stats['from']} to {stats['to']}"
    lo = stats["left_out"]
    overall = [
        {**stats["overall"], "what": f"Counted, {period}"},
        {**lo["excluded"], "what": "Left out: excluded readings"},
        {**lo["not_in_production"], "what": "Left out: not in production"},
        {**lo["excluded_stations"], "what": "Left out: stations excluded by default"},
    ]
    sheets = [
        ("Overall", ["Parts"], overall, ["what"]),
        ("Per station", ["Station"], stats["per_station"], ["station"]),
        ("Per job", ["Station", "Job"], stats["per_job"], ["station", "job"]),
        ("Per day", ["Day"], stats["per_day"], ["day"]),
    ]
    for i, (title, heads, rows, keys) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        marked = "station" in keys  # station rows say whether they are in the totals
        above = set((stats.get("alert") or {}).get("days_above") or []) if keys == ["day"] else None
        ws.append(heads + ["Pass", "Fail", "Total", "Scrap %"] + (["In totals"] if marked else [])
                  + (["Above scrap alert"] if above is not None and "alert" in stats else []))
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            day = [dt.date.fromisoformat(r["day"])] if keys == ["day"] else []
            ws.append(day + [r[k] for k in keys if k != "day"] + [r["pass"], r["fail"], r["total"], r["scrap_rate"]]
                      + (["No (excluded by default)" if r.get("excluded_by_default") else "Yes"] if marked else [])
                      + (["Yes" if r["day"] in above else ""] if above is not None and "alert" in stats else []))
        n = len(keys)
        for row in ws.iter_rows(min_row=2):
            row[n + 3].number_format = "0.00%"
            if keys == ["day"]:
                row[0].number_format = "yyyy-mm-dd"
        for col, width in zip("ABCDEFGH", [36 if keys == ["what"] else 24] * n + [12, 12, 12, 10, 24, 18]):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
